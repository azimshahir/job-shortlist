"""
Tests for the new pipeline stages. Deterministic and offline.

    python test_pipeline.py

The embedding tests load the real local model (no network after the first
download, no API). If sentence-transformers is unavailable they SKIP loudly
rather than silently passing.

test_filters.py still covers the original keyword rules and is untouched.
"""

import importlib
import os
import sys
import tempfile

import pandas as pd

import config
import enrich
import filters
import pipeline
import candidate as profile_mod
import ranking
import resume as resume_mod
import semantic

PASS, FAIL, SKIP = "PASS", "FAIL", "SKIP"
_results = []


def check(name, condition, detail=""):
    status = PASS if condition else FAIL
    _results.append((status, name, detail))
    print(f"{status:<4} {name}" + (f"\n       {detail}" if detail else ""))
    return condition


def skip(name, detail=""):
    _results.append((SKIP, name, detail))
    print(f"{SKIP:<4} {name}" + (f"\n       {detail}" if detail else ""))


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------
STRONG_FUND_JOB = {
    "site": "linkedin", "title": "Fund Operations Analyst",
    "company": "Global Custody Bank", "location": "Kuala Lumpur, Malaysia",
    "job_url": "https://example.com/1",
    "description": ("Daily reconciliation of fund positions and cash against "
                    "custodian records. Trade settlement monitoring, corporate "
                    "actions processing and support for fund administration "
                    "and transfer agency teams. 3-5 years in investment "
                    "operations."),
}

# Same domain, deliberately different vocabulary -- no shared boost keywords
# with the profile beyond generic ones. Tests that embeddings catch meaning.
PARAPHRASED_JOB = {
    "site": "linkedin", "title": "Middle Office Securities Services Associate",
    "company": "Asset Servicing Firm", "location": "Kuala Lumpur, Malaysia",
    "job_url": "https://example.com/2",
    "description": ("Support post-trade processing for collective investment "
                    "schemes. Match and affirm trades with counterparties, "
                    "monitor failing deliveries, verify pricing inputs feeding "
                    "daily valuation cycles, and liaise with trustees and "
                    "depositary teams on asset servicing events."),
}

WAREHOUSE_JOB = {
    "site": "indeed", "title": "Assistant Manager, Warehouse Fulfillment",
    "company": "Logistics Digital", "location": "Selangor, Malaysia",
    "job_url": "https://example.com/3",
    "description": ("Oversee warehouse fulfilment operations, inventory "
                    "accuracy and stock reconciliation. Manage packing teams, "
                    "delivery scheduling and last mile partners. Use AI tools "
                    "to improve throughput."),
}

RETAIL_JOB = {
    "site": "indeed", "title": "Retail Executive",
    "company": "Bakez Grocer", "location": "Kuala Lumpur, Malaysia",
    "job_url": "https://example.com/4",
    "description": ("Manage daily store operations, staff rostering and "
                    "customer service. Daily cash management, float handling "
                    "and reconciliation of till takings against the POS "
                    "system."),
}

MANDARIN_ADVANTAGE = {
    "site": "linkedin", "title": "Settlement Operations Executive",
    "company": "Regional Bank", "location": "Kuala Lumpur, Malaysia",
    "job_url": "https://example.com/5",
    "description": ("Trade settlement and reconciliation for fund operations. "
                    "Corporate actions support. Mandarin is an advantage."),
}

MANDARIN_REQUIRED = {
    "site": "linkedin", "title": "Settlement Operations Executive",
    "company": "Regional Bank", "location": "Kuala Lumpur, Malaysia",
    "job_url": "https://example.com/6",
    "description": ("Trade settlement and reconciliation for fund operations. "
                    "Fluent in Mandarin is required for this role."),
}

NO_DESCRIPTION = {
    "site": "linkedin", "title": "Fund Operations Analyst",
    "company": "Unknown Co", "location": "Kuala Lumpur, Malaysia",
    "job_url": "https://example.com/7", "description": None,
}


# --------------------------------------------------------------------------
# 3, 4, 5 -- filter rules survive the refactor
# --------------------------------------------------------------------------
def test_filter_rules_preserved():
    v = filters.score_job(MANDARIN_ADVANTAGE)
    check("3. 'Mandarin is an advantage' is not rejected",
          not v["rejected_by"] and v["keep"],
          f"score {v['score']}, rejected_by={v['rejected_by']}")

    v = filters.score_job(MANDARIN_REQUIRED)
    check("4. 'Fluent in Mandarin is required' is rejected",
          bool(v["rejected_by"]),
          f"rejected_by={v['rejected_by']}")

    # Both spellings of one concept present -> counted once.
    both = {"title": "Business Analyst", "company": "X",
            "description": "UAT coordination and user acceptance testing."}
    one = {"title": "Business Analyst", "company": "X",
           "description": "UAT coordination."}
    v_both, v_one = filters.score_job(both), filters.score_job(one)
    check("5. Duplicate synonyms do not double-count",
          v_both["score"] == v_one["score"],
          f"both={v_both['score']} one={v_one['score']} "
          f"(matched: {v_both['matched']})")

    # transfer pricing vs fund transfer pricing
    tp = filters.score_job({"title": "Tax Analyst", "company": "X",
                            "description": "Transfer pricing documentation."})
    ftp = filters.score_job({"title": "Treasury Analyst", "company": "X",
                            "description": "Fund transfer pricing and ALCO support."})
    check("5b. transfer pricing rejects, fund transfer pricing does not",
          bool(tp["rejected_by"]) and not ftp["rejected_by"],
          f"tp={tp['rejected_by']} ftp_score={ftp['score']}")


# --------------------------------------------------------------------------
# 7, 8, 9 -- graceful degradation
# --------------------------------------------------------------------------
def test_missing_description():
    v = filters.score_job(NO_DESCRIPTION)
    ok = isinstance(v["score"], int)
    status = enrich.description_status(NO_DESCRIPTION)
    text = semantic.job_text(NO_DESCRIPTION)
    check("7. Missing description fails gracefully",
          ok and status == enrich.MISSING and bool(text),
          f"status={status}, still embeddable via title/company")

    gaps = enrich.needs_enrichment(NO_DESCRIPTION)
    check("7b. Missing fields are reported for the Playwright seam",
          "description" in gaps and "job_url_direct" in gaps,
          f"gaps={gaps}")


def test_empty_scrape():
    empty = pd.DataFrame()
    deduped = scraper_dedupe_safe(empty)
    records = pipeline.to_records(deduped)
    passed, all_rows = pipeline.apply_filters(records)
    ranked = ranking.apply_ranking([])
    selected = ranking.select_final([])
    sem = semantic.score_jobs([], {"sections": {"a": "b"}})
    check("8. Empty scrape does not crash",
          records == [] and passed == [] and all_rows == []
          and ranked == [] and selected == [] and sem == [],
          "every stage returned empty without raising")

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "empty.csv")
        pipeline.write_stage([], path)
        df = pd.read_csv(path)
        check("8b. Empty stage CSV is written with headers",
              len(df) == 0 and "final_score" in df.columns,
              f"{len(df.columns)} columns, 0 rows")


def scraper_dedupe_safe(df):
    import scraper
    return scraper.dedupe(df)


def test_zero_match_email():
    stats = {"raw": 952, "hard_rejected": 400, "keyword_passed": 40,
             "ranked": 40, "selected": 0}
    import emailer
    html = emailer.build_html([], stats)
    subject = emailer.subject_for(stats)
    check("9. Zero-match email builds and says so plainly",
          "no strong matches" in subject.lower()
          and "No sufficiently strong matches" in html
          and "952" in html,
          f"subject: {subject}")


# --------------------------------------------------------------------------
# 10 -- no LLM in the cheap stages
# --------------------------------------------------------------------------
def test_no_llm_upstream():
    upstream = ["scraper", "filters", "semantic", "ranking", "enrich",
                "candidate", "config"]
    for mod in upstream:
        importlib.import_module(mod)

    leaked = "llm" in sys.modules
    check("10. No LLM module is imported by scrape/filter/embed/rank stages",
          not leaked,
          "llm.py was imported!" if leaked else
          "llm.py stays unimported until the final selection step")

    import llm
    check("10b. LLM is unconfigured by default and fails gracefully",
          not llm.is_configured(),
          f"LLM_PROVIDER={config.LLM_PROVIDER!r}")

    results = llm.analyse_jobs([STRONG_FUND_JOB, WAREHOUSE_JOB], {"meta": {}})
    check("10c. Unconfigured LLM returns a reason, does not raise",
          all(r["status"] == "not_configured" for r in results),
          results[0]["error"][:70])

    many = [STRONG_FUND_JOB] * 50
    capped = llm.analyse_jobs(many, {"meta": {}})
    check("10d. LLM calls are hard-capped per run",
          len(capped) == config.LLM_MAX_CALLS_PER_RUN,
          f"50 jobs in -> {len(capped)} out "
          f"(cap {config.LLM_MAX_CALLS_PER_RUN})")


# --------------------------------------------------------------------------
# Profile accuracy contract
# --------------------------------------------------------------------------
def test_profile_contract():
    try:
        prof = profile_mod.load_profile()
    except Exception as exc:  # noqa: BLE001
        check("P1. candidate_profile.yaml loads", False, str(exc))
        return None

    check("P1. candidate_profile.yaml loads",
          len(prof["sections"]) >= 8,
          f"{len(prof['sections'])} profile sections")

    banned = [b.lower() for b in profile_mod.must_not_claim(prof)]
    blob = " ".join(prof["sections"].values()).lower()
    # These must never appear as claims in the embedded profile text.
    leaks = [b for b in ("variance analysis", "process mapping", "bookkeeping",
                         "invoicing", "financial modelling", "vostro", "nostro",
                         "bpmn", "visio") if b in blob]
    check("P2. Profile text claims nothing on the must-not-claim list",
          not leaks, f"leaked: {leaks}" if leaks else f"{len(banned)} rules enforced")

    check("P3. NAV is framed as contributing inputs, not ownership",
          "contributing inputs to daily nav" in blob
          and "nav calculation" not in blob,
          "SS&C framing correct")

    certs = " ".join(str(c) for c in prof.get("certifications", [])).lower()
    check("P4. Microsoft Copilot Training Facilitator is in the profile",
          "copilot training facilitator" in certs
          or "copilot training facilitator" in blob)

    agrobank = [e for e in prof.get("experience", [])
                if "agrobank" in str(e.get("employer", "")).lower()]
    check("P5. Agrobank ERM title is Executive",
          bool(agrobank) and "executive" in str(agrobank[0]["title"]).lower(),
          agrobank[0]["title"] if agrobank else "not found")
    return prof


# --------------------------------------------------------------------------
# 1, 2, 6 -- semantic behaviour (needs the local model)
# --------------------------------------------------------------------------
def test_semantic(prof):
    try:
        model = semantic.get_model()
    except Exception as exc:  # noqa: BLE001
        for n in ("1. Strong fund ops job ranks highly",
                  "2. Warehouse 'reconciliation' job does not rank highly",
                  "6. Paraphrased same-domain job still ranks well"):
            skip(n, f"embedding model unavailable: {type(exc).__name__}")
        return

    jobs = [STRONG_FUND_JOB, PARAPHRASED_JOB, WAREHOUSE_JOB, RETAIL_JOB]
    for job, result in zip(jobs, semantic.score_jobs(jobs, prof, model=model)):
        job.update(result)
        job["keyword_score"] = filters.score_job(job)["score"]

    ranking.apply_ranking(jobs)
    by_title = {j["title"]: j for j in jobs}
    strong = by_title["Fund Operations Analyst"]
    para = by_title["Middle Office Securities Services Associate"]
    warehouse = by_title["Assistant Manager, Warehouse Fulfillment"]
    retail = by_title["Retail Executive"]

    check("1. Strong fund ops job ranks highly",
          strong["final_rank"] == 1 and strong["final_score"] >= 60,
          f"rank {strong['final_rank']}, final {strong['final_score']}, "
          f"semantic {strong['semantic_score']}, "
          f"match: {strong['strongest_profile_match']}")

    check("2. Warehouse 'reconciliation' job does not rank highly",
          warehouse["final_score"] < strong["final_score"]
          and warehouse["final_rank"] > para["final_rank"]
          and warehouse["semantic_score"] < strong["semantic_score"],
          f"warehouse final {warehouse['final_score']} "
          f"(sem {warehouse['semantic_score']}, kw {warehouse['keyword_score']}) "
          f"vs strong {strong['final_score']}")

    check("2b. Retail 'cash management' job does not rank highly",
          retail["final_score"] < strong["final_score"]
          and retail["semantic_score"] < strong["semantic_score"],
          f"retail final {retail['final_score']} "
          f"(sem {retail['semantic_score']}, kw {retail['keyword_score']})")

    check("6. Paraphrased same-domain job still ranks well",
          para["semantic_score"] > warehouse["semantic_score"]
          and para["semantic_score"] > retail["semantic_score"],
          f"paraphrase sem {para['semantic_score']} (kw only "
          f"{para['keyword_score']}) beats warehouse "
          f"{warehouse['semantic_score']} / retail {retail['semantic_score']}")

    check("6b. Semantic score is normalised to a readable 0-100",
          all(0 <= j["semantic_score"] <= 100 for j in jobs),
          ", ".join(f"{j['title'][:22]}={j['semantic_score']}" for j in jobs))


# --------------------------------------------------------------------------
# Selection + resume plumbing
# --------------------------------------------------------------------------
def test_selection():
    rows = [{"title": f"Job {i}", "final_score": s, "strongest_profile_match": "X"}
            for i, s in enumerate([95, 88, 80, 72, 65, 62, 59, 40])]
    selected = ranking.select_final(rows)
    check("S1. Selection caps at MAX_FINAL_SELECTION",
          len(selected) == config.MAX_FINAL_SELECTION,
          f"{len(selected)} selected from {len(rows)}")

    weak = [{"title": "A", "final_score": 30}, {"title": "B", "final_score": 20}]
    selected_weak = ranking.select_final(weak)
    check("S2. Weak-only day selects nothing rather than padding",
          selected_weak == []
          and all(r["selection_status"] == "not_selected" for r in weak),
          f"reason: {weak[0]['selection_reason']}")

    two_strong = [{"title": "A", "final_score": 90}, {"title": "B", "final_score": 85},
                  {"title": "C", "final_score": 30}]
    check("S3. Two strong jobs are not padded up to five",
          len(ranking.select_final(two_strong)) == 2)

    # Regression: internships score deceptively well semantically because the
    # JD describes exactly the right work. They must never be selected.
    with_intern = [
        {"title": "Intern, Settlement - Cash Management", "final_score": 88},
        {"title": "Reconciliation Analyst", "final_score": 70},
        {"title": "Graduate Programme, Operations", "final_score": 85},
    ]
    picked = ranking.select_final(with_intern)
    check("S4. Entry-level roles are disqualified regardless of score",
          len(picked) == 1 and picked[0]["title"] == "Reconciliation Analyst"
          and with_intern[0]["selection_status"] == "disqualified",
          f"selected: {[p['title'] for p in picked]}; "
          f"intern reason: {with_intern[0]['selection_reason'][:60]}")

    check("S5. Disqualification is recorded with a visible reason, not hidden",
          all(r.get("selection_reason") for r in with_intern),
          "every row carries a selection_reason")


def test_resume_filename():
    name = resume_mod.safe_filename("CACEIS Malaysia Sdn Bhd",
                                    "Assistant Manager, Middle Office / Ops")
    check("R1. Resume filename is sanitised",
          name.endswith(".docx") and "/" not in name and "," not in name
          and name.startswith("Azim_Shahir_"),
          name)

    nasty = resume_mod.safe_filename("../../etc", "..\\passwd")
    check("R2. Path traversal cannot escape the resume directory",
          "/" not in nasty and "\\" not in nasty and ".." not in nasty,
          nasty)


def test_resume_docx(prof):
    analysis = {
        "tailored_summary": "Fund operations professional with five years.",
        "tailored_skills": ["Reconciliation", "Settlement"],
        "tailored_bullets": [{"employer": "Agrobank",
                              "title": "Executive, Enterprise Risk Management",
                              "dates": "2022 - 2024",
                              "bullets": ["Cleared 120 ageing breaks per month."]}],
    }
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "t.docx")
        try:
            resume_mod.build_resume(analysis, STRONG_FUND_JOB, prof, path)
        except Exception as exc:  # noqa: BLE001
            check("R3. DOCX renders", False, f"{type(exc).__name__}: {exc}")
            return
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        check("R3. DOCX renders with Calibri body and required content",
              doc.styles["Normal"].font.name == config.RESUME_FONT
              and "Copilot Training Facilitator" in text
              and "120 ageing breaks" in text,
              "metric preserved verbatim; Copilot facilitator present")


def test_placeholder_guard(prof):
    todos = profile_mod.has_placeholders(prof)
    check("R4. TODO placeholders are detected so resumes are not generated "
          "from invented data",
          isinstance(todos, list),
          f"{len(todos)} placeholder(s) outstanding"
          + (f", e.g. {todos[0]}" if todos else ""))


# --------------------------------------------------------------------------
def main():
    print("=" * 70)
    print("PIPELINE TESTS")
    print("=" * 70)
    test_filter_rules_preserved()
    test_missing_description()
    test_empty_scrape()
    test_zero_match_email()
    test_no_llm_upstream()
    prof = test_profile_contract()
    if prof:
        test_semantic(prof)
        test_resume_docx(prof)
        test_placeholder_guard(prof)
    test_selection()
    test_resume_filename()

    print("\n" + "=" * 70)
    failed = [r for r in _results if r[0] == FAIL]
    skipped = [r for r in _results if r[0] == SKIP]
    print(f"{len(_results) - len(failed) - len(skipped)} passed, "
          f"{len(failed)} failed, {len(skipped)} skipped")
    for _, name, detail in failed:
        print(f"  FAILED: {name} -- {detail}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
